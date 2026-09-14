from datetime import datetime, timedelta
from io import BytesIO
import re

from openpyxl import Workbook
from flask import Blueprint, flash, redirect, render_template, request, send_file, session, url_for
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from db import db
import admin_phone_numbers_legacy as legacy
from phone_history import verification_settings

admin_phone_numbers_bp = Blueprint("admin_phone_numbers", __name__)

orders_col = db["orders"]
blocked_phone_numbers_col = db["blocked_phone_numbers"]
not_in_database_phones_col = db["not_in_database_phone_numbers"]
phone_verification_settings_col = db["phone_verification_settings"]
services_col = db["services"]
PHONE_VERIFICATION_SETTINGS_ID = "PHONE_HISTORY_SETTINGS"

EXPORT_NETWORK_KEYWORDS = {
    "MTN": ("mtn",),
    "TELECEL": ("telecel", "vodafone"),
    "AIRTELTIGO": ("airteltigo", "airtel", "tigo", "at"),
}


def _require_admin():
    return session.get("role") == "main_admin" and bool(session.get("user_id"))


def _normalize_phone(phone: str) -> str:
    digits = re.sub(r"\D+", "", str(phone or ""))
    if not digits:
        return ""
    if digits.startswith("233") and len(digits) == 12:
        return f"0{digits[3:]}"
    if len(digits) == 9:
        return f"0{digits}"
    return digits


def _pagination_params():
    page_raw = request.args.get("page", 1)
    try:
        page = max(int(page_raw), 1)
    except Exception:
        page = 1
    per_page = 50
    return page, per_page


def _date_params():
    start_date = (request.args.get("start_date") or request.form.get("start_date") or "").strip()
    end_date = (request.args.get("end_date") or request.form.get("end_date") or "").strip()
    return start_date, end_date


def _clean_text(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _normalize_network(value: str | None) -> str:
    raw = _clean_text(value).upper()
    if not raw or raw == "ALL":
        return ""
    for label, needles in EXPORT_NETWORK_KEYWORDS.items():
        if raw == label:
            return label
        for needle in needles:
            if needle.upper() in raw:
                return label
    return raw


def _parse_date_range(start_date: str, end_date: str) -> tuple[datetime | None, datetime | None]:
    start_dt = None
    end_dt = None
    try:
        if start_date:
            start_dt = datetime.strptime(start_date, "%Y-%m-%d")
    except Exception:
        start_dt = None
    try:
        if end_date:
            end_dt = datetime.strptime(end_date, "%Y-%m-%d") + timedelta(days=1)
    except Exception:
        end_dt = None
    return start_dt, end_dt


def _created_at_match(start_date: str, end_date: str) -> dict:
    start_dt, end_dt = _parse_date_range(start_date, end_date)
    date_match = {}
    if start_dt is not None:
        date_match["$gte"] = start_dt
    if end_dt is not None:
        date_match["$lt"] = end_dt
    return {"created_at": date_match} if date_match else {}


def _range_label(start_date: str, end_date: str) -> str:
    if start_date and end_date:
        return f"{start_date} to {end_date}"
    if start_date:
        return f"From {start_date}"
    if end_date:
        return f"Until {end_date}"
    return ""


def _build_network_patterns(network: str) -> list[str]:
    normalized = _normalize_network(network)
    if not normalized:
        return []
    patterns = [normalized]
    patterns.extend(EXPORT_NETWORK_KEYWORDS.get(normalized, ()))
    return sorted({p for p in patterns if p}, key=len, reverse=True)


def _service_network_options() -> list[str]:
    options = set(EXPORT_NETWORK_KEYWORDS.keys())
    try:
        cursor = services_col.find({}, {"name": 1, "service_network": 1, "network": 1})
        for doc in cursor:
            for key in ("service_network", "network", "name"):
                normalized = _normalize_network(doc.get(key))
                if normalized:
                    options.add(normalized)
    except Exception:
        pass
    return sorted(options)


def _existing_mtn_history_enforced() -> bool:
    return verification_settings()["require_existing_mtn_history"]


def _first_time_alert_settings() -> dict:
    default_message = "This is a New number, order delivery may take 24hrs"
    try:
        doc = phone_verification_settings_col.find_one(
            {"_id": PHONE_VERIFICATION_SETTINGS_ID},
            {"first_time_alert_enabled": 1, "first_time_alert_message": 1},
        ) or {}
        return {
            "enabled": bool(doc.get("first_time_alert_enabled", False)),
            "message": str(doc.get("first_time_alert_message") or default_message).strip(),
        }
    except Exception:
        return {"enabled": False, "message": default_message}


def _item_network_match(network: str) -> dict:
    patterns = _build_network_patterns(network)
    if not patterns:
        return {}
    regex = {"$regex": "|".join(re.escape(pattern) for pattern in patterns), "$options": "i"}
    return {
        "$or": [
            {"items.provider_network": regex},
            {"items.network": regex},
            {"items.network_name": regex},
            {"items.ported_expected_network": regex},
            {"items.ported_detected_network": regex},
            {"items.serviceName": regex},
            {"items.service_name": regex},
        ]
    }


def _fetch_order_phone_rows(
    q: str,
    network: str = "",
    start_date: str = "",
    end_date: str = "",
    skip: int | None = None,
    limit: int | None = None,
):
    if q:
        phone_match = {"$regex": re.escape(q), "$options": "i"}
    else:
        phone_match = {"$exists": True, "$nin": [None, ""]}

    base_match = {"items.phone": phone_match}
    base_match.update(_item_network_match(network))
    created_match = _created_at_match(start_date, end_date)

    total_pipeline = [
        *([{"$match": created_match}] if created_match else []),
        {"$unwind": "$items"},
        {"$match": base_match},
        {
            "$project": {
                "phone": "$items.phone",
                "provider_network": "$items.provider_network",
                "network": "$items.network",
                "network_name": "$items.network_name",
                "ported_expected_network": "$items.ported_expected_network",
                "ported_detected_network": "$items.ported_detected_network",
                "serviceName": "$items.serviceName",
                "service_name": "$items.service_name",
            }
        },
    ]
    total_rows = list(orders_col.aggregate(total_pipeline))
    total = len({_normalize_phone(row.get("phone")) for row in total_rows if re.fullmatch(r"0\d{9}", _normalize_phone(row.get("phone")))})

    # Count-only mode: avoid building a data pipeline with invalid/zero limits.
    if limit is not None and int(limit) <= 0:
        return total, []

    pipeline = [
        *([{"$match": created_match}] if created_match else []),
        {"$unwind": "$items"},
        {"$match": base_match},
        {
            "$project": {
                "order_id": "$order_id",
                "created_at": "$created_at",
                "phone": "$items.phone",
                "provider_network": "$items.provider_network",
                "network": "$items.network",
                "network_name": "$items.network_name",
                "ported_expected_network": "$items.ported_expected_network",
                "ported_detected_network": "$items.ported_detected_network",
                "serviceName": "$items.serviceName",
                "service_name": "$items.service_name",
            }
        },
    ]
    rows = list(orders_col.aggregate(pipeline))
    grouped_rows: dict[str, dict] = {}
    for row in rows:
        phone = row.get("phone")
        key = _normalize_phone(phone)
        if not re.fullmatch(r"0\d{9}", key):
            continue
        current = grouped_rows.get(key)
        created_at = row.get("created_at")
        if not current:
            grouped_rows[key] = {
                "phone": phone,
                "orders_count": 1 if row.get("order_id") else 0,
                "order_ids": {row.get("order_id")} if row.get("order_id") else set(),
                "last_order_at": created_at,
            }
            continue
        if row.get("order_id") and row.get("order_id") not in current["order_ids"]:
            current["order_ids"].add(row.get("order_id"))
        current["orders_count"] = len(current["order_ids"])
        if created_at and (
            not current["last_order_at"] or created_at > current["last_order_at"]
        ):
            current["last_order_at"] = created_at

    grouped_list = []
    for item in grouped_rows.values():
        item.pop("order_ids", None)
        grouped_list.append(item)

    grouped_list.sort(
        key=lambda item: (
            int(item.get("orders_count") or 0),
            str(item.get("phone") or ""),
            item.get("last_order_at") or datetime.min,
        )
    )

    if skip is not None:
        grouped_list = grouped_list[skip:]
    if limit is not None and int(limit) > 0:
        grouped_list = grouped_list[:limit]

    row_keys = [_normalize_phone(r.get("phone")) for r in grouped_list if r.get("phone")]
    active_blocks = list(
        blocked_phone_numbers_col.find(
            {"is_active": True, "normalized_phone": {"$in": row_keys}},
            {"normalized_phone": 1, "reason": 1, "_id": 0},
        )
    )
    blocked_map = {d.get("normalized_phone"): d for d in active_blocks if d.get("normalized_phone")}

    for row in grouped_list:
        key = _normalize_phone(row.get("phone"))
        row["normalized_phone"] = key
        row["is_blocked"] = key in blocked_map
        row["block_reason"] = (blocked_map.get(key) or {}).get("reason", "")

    return total, grouped_list


def _phone_variants(normalized_phone: str) -> list[str]:
    key = _normalize_phone(normalized_phone)
    if not key:
        return []

    variants = {key}
    digits = re.sub(r"\D+", "", key)
    if digits.startswith("0") and len(digits) == 10:
        variants.add(digits[1:])
        variants.add(f"233{digits[1:]}")
    elif digits.startswith("233") and len(digits) == 12:
        variants.add(f"0{digits[3:]}")
        variants.add(digits[3:])
    elif len(digits) == 9:
        variants.add(f"0{digits}")
        variants.add(f"233{digits}")

    return list(variants)


def _fetch_order_stats_for_phones(
    normalized_phones: list[str],
    network: str = "",
    start_date: str = "",
    end_date: str = "",
) -> dict[str, dict]:
    variants = []
    for phone in normalized_phones:
        variants.extend(_phone_variants(phone))

    variants = sorted({v for v in variants if v})
    if not variants:
        return {}

    created_match = _created_at_match(start_date, end_date)
    pipeline = [
        *([{"$match": created_match}] if created_match else []),
        {"$unwind": "$items"},
        {"$match": {"items.phone": {"$in": variants}, **_item_network_match(network)}},
        {
            "$project": {
                "order_id": "$order_id",
                "created_at": "$created_at",
                "phone": "$items.phone",
                "provider_network": "$items.provider_network",
                "network": "$items.network",
                "network_name": "$items.network_name",
                "ported_expected_network": "$items.ported_expected_network",
                "ported_detected_network": "$items.ported_detected_network",
                "serviceName": "$items.serviceName",
                "service_name": "$items.service_name",
            }
        },
    ]
    raw_rows = list(orders_col.aggregate(pipeline))

    stats: dict[str, dict] = {}
    for row in raw_rows:
        key = _normalize_phone(row.get("phone"))
        if not key:
            continue
        current = stats.get(key) or {"orders_count": 0, "last_order_at": None}
        current_order_ids = current.setdefault("order_ids", set())
        if row.get("order_id"):
            current_order_ids.add(row.get("order_id"))
        current["orders_count"] = len(current_order_ids)
        last_order_at = row.get("created_at")
        if last_order_at and (
            not current["last_order_at"] or last_order_at > current["last_order_at"]
        ):
            current["last_order_at"] = last_order_at
        stats[key] = current

    for item in stats.values():
        item.pop("order_ids", None)

    return stats


def _fetch_blocked_phone_rows(
    q: str,
    network: str = "",
    start_date: str = "",
    end_date: str = "",
    skip: int | None = None,
    limit: int | None = None,
):
    filters = {"is_active": True}
    if q:
        filters["$or"] = [
            {"phone": {"$regex": re.escape(q), "$options": "i"}},
            {"normalized_phone": {"$regex": re.escape(q), "$options": "i"}},
        ]

    cursor = blocked_phone_numbers_col.find(
        filters,
        {
            "_id": 0,
            "phone": 1,
            "normalized_phone": 1,
            "reason": 1,
            "created_at": 1,
            "updated_at": 1,
        },
    ).sort([("updated_at", -1), ("phone", 1)])

    rows = list(cursor)
    stats_map = _fetch_order_stats_for_phones(
        [row.get("normalized_phone") for row in rows if row.get("normalized_phone")],
        network=network,
        start_date=start_date,
        end_date=end_date,
    )

    normalized_q = _normalize_phone(q) if q else ""
    if normalized_q:
        rows = [
            row
            for row in rows
            if normalized_q in (row.get("normalized_phone") or "")
            or normalized_q in _normalize_phone(row.get("phone"))
        ]

    filtered_rows = []
    for row in rows:
        key = row.get("normalized_phone") or _normalize_phone(row.get("phone"))
        stats = stats_map.get(key) or {}
        row["phone"] = row.get("phone") or key
        row["normalized_phone"] = key
        row["orders_count"] = int(stats.get("orders_count") or 0)
        row["last_order_at"] = stats.get("last_order_at")
        row["is_blocked"] = True
        row["block_reason"] = row.get("reason") or ""
        if not network and not start_date and not end_date:
            filtered_rows.append(row)
        elif row["orders_count"] > 0:
            filtered_rows.append(row)

    total = len(filtered_rows)
    if limit is not None and int(limit) <= 0:
        return total, []
    if skip is not None:
        filtered_rows = filtered_rows[skip:]
    if limit is not None and int(limit) > 0:
        filtered_rows = filtered_rows[:limit]

    return total, filtered_rows


def _fetch_not_in_database_rows(
    q: str,
    network: str = "",
    start_date: str = "",
    end_date: str = "",
    skip: int | None = None,
    limit: int | None = None,
):
    filters = {}
    if q:
        filters["$or"] = [
            {"phone": {"$regex": re.escape(q), "$options": "i"}},
            {"normalized_phone": {"$regex": re.escape(q), "$options": "i"}},
            {"last_service_name": {"$regex": re.escape(q), "$options": "i"}},
        ]
    if network:
        filters["network_label"] = _normalize_network(network)
    last_seen_match = _created_at_match(start_date, end_date)
    if last_seen_match:
        filters["last_seen_at"] = last_seen_match["created_at"]

    cursor = not_in_database_phones_col.find(
        filters,
        {
            "_id": 0,
            "phone": 1,
            "normalized_phone": 1,
            "first_seen_at": 1,
            "last_seen_at": 1,
            "seen_count": 1,
            "last_source": 1,
            "last_service_name": 1,
            "last_service_network": 1,
            "last_network": 1,
            "network_label": 1,
            "allow_order_override": 1,
            "approved_at": 1,
            "approved_by": 1,
        },
    ).sort([("last_seen_at", -1), ("phone", 1)])

    rows = list(cursor)
    stats_map = _fetch_order_stats_for_phones(
        [row.get("normalized_phone") or _normalize_phone(row.get("phone")) for row in rows],
        network=network,
        start_date=start_date,
        end_date=end_date,
    )

    for row in rows:
        key = row.get("normalized_phone") or _normalize_phone(row.get("phone"))
        stats = stats_map.get(key) or {}
        row["phone"] = row.get("phone") or key
        row["normalized_phone"] = key
        row["orders_count"] = int(stats.get("orders_count") or 0)
        row["last_order_at"] = stats.get("last_order_at")
        row["is_blocked"] = False
        row["block_reason"] = ""
        row["allow_order_override"] = bool(row.get("allow_order_override"))

    total = len(rows)
    if limit is not None and int(limit) <= 0:
        return total, []
    if skip is not None:
        rows = rows[skip:]
    if limit is not None and int(limit) > 0:
        rows = rows[:limit]
    return total, rows


def _fetch_phone_rows(
    q: str,
    status: str = "all",
    network: str = "",
    start_date: str = "",
    end_date: str = "",
    skip: int | None = None,
    limit: int | None = None,
):
    if status == "blocked":
        return _fetch_blocked_phone_rows(q=q, network=network, start_date=start_date, end_date=end_date, skip=skip, limit=limit)
    if status == "not_in_database":
        return _fetch_not_in_database_rows(q=q, network=network, start_date=start_date, end_date=end_date, skip=skip, limit=limit)
    return _fetch_order_phone_rows(q=q, network=network, start_date=start_date, end_date=end_date, skip=skip, limit=limit)


@admin_phone_numbers_bp.route("/admin/phone-numbers")
def phone_numbers_page():
    if session.get("role") == "admin":
        return legacy.phone_numbers_page()
    if not _require_admin():
        return redirect(url_for("login.login"))

    q = (request.args.get("q") or "").strip()
    status = (request.args.get("status") or "all").strip().lower()
    network = _normalize_network(request.args.get("network"))
    start_date, end_date = _date_params()
    if status not in {"all", "blocked", "not_in_database"}:
        status = "all"
    page, per_page = _pagination_params()

    total, _ = _fetch_phone_rows(q=q, status=status, network=network, start_date=start_date, end_date=end_date, skip=0, limit=0)
    total_pages = max((total + per_page - 1) // per_page, 1)
    if page > total_pages:
        page = total_pages
    skip = (page - 1) * per_page

    _, rows = _fetch_phone_rows(q=q, status=status, network=network, start_date=start_date, end_date=end_date, skip=skip, limit=per_page)

    total_blocked = blocked_phone_numbers_col.count_documents({"is_active": True})
    total_not_in_database = not_in_database_phones_col.count_documents({})

    first_time_alert = _first_time_alert_settings()
    return render_template(
        "admin_phone_numbers.html",
        rows=rows,
        q=q,
        status=status,
        network=network,
        start_date=start_date,
        end_date=end_date,
        date_range_label=_range_label(start_date, end_date),
        network_options=_service_network_options(),
        page=page,
        per_page=per_page,
        total=total,
        total_pages=total_pages,
        total_blocked=total_blocked,
        total_not_in_database=total_not_in_database,
        require_existing_mtn_history=_existing_mtn_history_enforced(),
        first_time_alert_enabled=first_time_alert["enabled"],
        first_time_alert_message=first_time_alert["message"],
    )


@admin_phone_numbers_bp.route("/admin/phone-numbers/export/excel")
def export_phone_numbers_excel():
    if session.get("role") == "admin":
        return legacy.export_phone_numbers_excel()
    if not _require_admin():
        return redirect(url_for("login.login"))

    q = (request.args.get("q") or "").strip()
    status = (request.args.get("status") or "all").strip().lower()
    network = _normalize_network(request.args.get("network"))
    start_date, end_date = _date_params()
    if status not in {"all", "blocked", "not_in_database"}:
        status = "all"
    _, rows = _fetch_phone_rows(q=q, status=status, network=network, start_date=start_date, end_date=end_date)

    if status == "not_in_database":
        grouped: dict[str, list[str]] = {}
        for row in rows:
            phone = _normalize_phone(row.get("phone"))
            if not phone:
                continue
            network_label = _normalize_network(row.get("network_label")) or "OTHER"
            grouped.setdefault(network_label, []).append(phone)

        preferred_order = ["MTN", "TELECEL", "AIRTELTIGO", "OTHER"]
        headers = [name for name in preferred_order if grouped.get(name)]
        headers.extend(sorted(name for name in grouped if name not in preferred_order))
        columns = [sorted(set(grouped[header])) for header in headers]
        sheet_name = "Not in Database"
        download_name = "not_in_database_numbers.xlsx"
    else:
        column_name = network if network else "Phone Numbers"
        headers = [column_name]
        columns = [[row.get("phone", "") for row in rows]]
        sheet_name = (network or "Phone Numbers")[:31]
        filename_prefix = re.sub(r"[^A-Za-z0-9]+", "_", (network or "phone_numbers")).strip("_") or "phone_numbers"
        download_name = f"{filename_prefix}_numbers.xlsx"

    output = BytesIO()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(headers or ["Phone Numbers"])
    for index in range(max((len(column) for column in columns), default=0)):
        worksheet.append([column[index] if index < len(column) else "" for column in columns])
    worksheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="0F172A")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
        worksheet.column_dimensions[cell.column_letter].width = 20
    workbook.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=download_name,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@admin_phone_numbers_bp.route("/admin/phone-numbers/export/pdf")
def export_phone_numbers_pdf():
    if session.get("role") == "admin":
        return legacy.export_phone_numbers_pdf()
    if not _require_admin():
        return redirect(url_for("login.login"))

    q = (request.args.get("q") or "").strip()
    status = (request.args.get("status") or "all").strip().lower()
    network = _normalize_network(request.args.get("network"))
    start_date, end_date = _date_params()
    if status not in {"all", "blocked", "not_in_database"}:
        status = "all"
    total, rows = _fetch_phone_rows(q=q, status=status, network=network, start_date=start_date, end_date=end_date)

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=24,
        leftMargin=24,
        topMargin=24,
        bottomMargin=24,
    )

    styles = getSampleStyleSheet()
    title = Paragraph("Phone Numbers Report", styles["Title"])
    subtitle = Paragraph(
        f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')} | Total: {total}"
        + (f" | Network: {network}" if network else ""),
        styles["Normal"],
    )
    if start_date or end_date:
        subtitle = Paragraph(
            f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')} | Total: {total}"
            + (f" | Network: {network}" if network else "")
            + f" | Date: {_range_label(start_date, end_date)}",
            styles["Normal"],
        )

    table_data = [["#", "Phone Number", "Orders", "Status", "Reason", "Last Order"]]
    for idx, row in enumerate(rows, start=1):
        table_data.append(
            [
                str(idx),
                str(row.get("phone") or ""),
                str(int(row.get("orders_count") or 0)),
                "Blocked" if row.get("is_blocked") else "Active",
                str(row.get("block_reason") or ""),
                row.get("last_order_at").strftime("%Y-%m-%d %H:%M")
                if row.get("last_order_at")
                else "-",
            ]
        )

    tbl = Table(table_data, repeatRows=1, colWidths=[36, 140, 70, 80, 260, 110])
    tbl.setStyle(
        TableStyle(
            [
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#0f172a")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                ("ALIGN", (0, 0), (0, -1), "CENTER"),
                ("ALIGN", (2, 0), (3, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )

    doc.build([title, Spacer(1, 8), subtitle, Spacer(1, 12), tbl])
    buffer.seek(0)

    stamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    return send_file(
        buffer,
        as_attachment=True,
        download_name=f"phone_numbers_{stamp}.pdf",
        mimetype="application/pdf",
    )


@admin_phone_numbers_bp.route("/admin/phone-numbers/verification-setting", methods=["POST"])
def update_phone_verification_setting():
    if not _require_admin():
        return redirect(url_for("login.login"))

    enabled = (request.form.get("require_existing_mtn_history") or "").strip().lower() == "on"
    now = datetime.utcnow()
    phone_verification_settings_col.update_one(
        {"_id": PHONE_VERIFICATION_SETTINGS_ID},
        {
            "$set": {
                "require_existing_mtn_history": enabled,
                "updated_at": now,
                "updated_by": session.get("admin_id") or session.get("user_id"),
            },
            "$setOnInsert": {"created_at": now},
        },
        upsert=True,
    )
    flash(
        (
            "MTN phone-history verification is ON."
            if enabled
            else "MTN phone-history verification is OFF. All numbers can order without verification or warnings."
        ),
        "success",
    )
    return redirect(url_for("admin_phone_numbers.phone_numbers_page", status="not_in_database"))


@admin_phone_numbers_bp.route("/admin/phone-numbers/first-time-alert-setting", methods=["POST"])
def update_first_time_alert_setting():
    if not _require_admin():
        return redirect(url_for("login.login"))

    enabled = (request.form.get("first_time_alert_enabled") or "").strip().lower() == "on"
    message = re.sub(r"\s+", " ", request.form.get("first_time_alert_message") or "").strip()
    if not message:
        message = "This is a New number, order delivery may take 24hrs"
    message = message[:300]
    now = datetime.utcnow()
    phone_verification_settings_col.update_one(
        {"_id": PHONE_VERIFICATION_SETTINGS_ID},
        {
            "$set": {
                "first_time_alert_enabled": enabled,
                "first_time_alert_message": message,
                "updated_at": now,
                "updated_by": session.get("admin_id") or session.get("user_id"),
            },
            "$setOnInsert": {"created_at": now},
        },
        upsert=True,
    )
    flash(f"First-Time Alert is {'ON' if enabled else 'OFF'}.", "success")
    return redirect(url_for("admin_phone_numbers.phone_numbers_page"))


@admin_phone_numbers_bp.route("/admin/phone-numbers/not-in-database/unblock", methods=["POST"])
def unblock_not_in_database_phone():
    if not _require_admin():
        return redirect(url_for("login.login"))

    phone = (request.form.get("phone") or "").strip()
    q = (request.form.get("q") or "").strip()
    status = (request.form.get("status") or "not_in_database").strip().lower()
    network = _normalize_network(request.form.get("network"))
    page = (request.form.get("page") or "1").strip()
    start_date = (request.form.get("start_date") or "").strip()
    end_date = (request.form.get("end_date") or "").strip()

    key = _normalize_phone(phone)
    if re.fullmatch(r"0\d{9}", key):
        now = datetime.utcnow()
        not_in_database_phones_col.update_one(
            {"normalized_phone": key},
            {
                "$set": {
                    "phone": phone or key,
                    "normalized_phone": key,
                    "allow_order_override": True,
                    "approved_at": now,
                    "approved_by": session.get("admin_id") or session.get("user_id"),
                    "updated_at": now,
                },
                "$setOnInsert": {
                    "first_seen_at": now,
                    "last_seen_at": now,
                    "seen_count": 0,
                },
            },
            upsert=True,
        )
        flash(f"{key} can now place orders even if it is not yet in the database.", "success")

    return redirect(
        url_for(
            "admin_phone_numbers.phone_numbers_page",
            q=q,
            status=status,
            network=network,
            start_date=start_date,
            end_date=end_date,
            page=page,
        )
    )


@admin_phone_numbers_bp.route("/admin/phone-numbers/bulk-unblock", methods=["POST"])
def bulk_unblock_phone_numbers():
    if not _require_admin():
        return redirect(url_for("login.login"))

    q = (request.form.get("q") or "").strip()
    status = (request.form.get("status") or "").strip().lower()
    network = _normalize_network(request.form.get("network"))
    page = (request.form.get("page") or "1").strip()
    start_date = (request.form.get("start_date") or "").strip()
    end_date = (request.form.get("end_date") or "").strip()
    scope = (request.form.get("scope") or "selected").strip().lower()

    if status not in {"blocked", "not_in_database"}:
        flash("Bulk unblock is only available for Blocked and Not in Database numbers.", "warning")
        return redirect(url_for("admin_phone_numbers.phone_numbers_page", status="all"))

    if scope == "all_filtered":
        _, candidate_rows = _fetch_phone_rows(
            q=q,
            status=status,
            network=network,
            start_date=start_date,
            end_date=end_date,
        )
        phones = [row.get("phone") for row in candidate_rows]
    else:
        phones = request.form.getlist("phones")

    keys = sorted({_normalize_phone(phone) for phone in phones if re.fullmatch(r"0\d{9}", _normalize_phone(phone))})
    if not keys:
        flash("Select at least one phone number to unblock.", "warning")
    else:
        now = datetime.utcnow()
        actor_id = session.get("admin_id") or session.get("user_id")
        if status == "blocked":
            result = blocked_phone_numbers_col.update_many(
                {"normalized_phone": {"$in": keys}, "is_active": True},
                {
                    "$set": {
                        "is_active": False,
                        "updated_at": now,
                        "unblocked_by": actor_id,
                    }
                },
            )
        else:
            result = not_in_database_phones_col.update_many(
                {
                    "normalized_phone": {"$in": keys},
                    "allow_order_override": {"$ne": True},
                },
                {
                    "$set": {
                        "allow_order_override": True,
                        "approved_at": now,
                        "approved_by": actor_id,
                        "updated_at": now,
                    }
                },
            )
        flash(f"{result.modified_count} phone number(s) unblocked successfully.", "success")

    return redirect(
        url_for(
            "admin_phone_numbers.phone_numbers_page",
            q=q,
            status=status,
            network=network,
            start_date=start_date,
            end_date=end_date,
            page=page,
        )
    )


@admin_phone_numbers_bp.route("/admin/phone-numbers/block", methods=["POST"])
def block_phone_number():
    if session.get("role") == "admin":
        return legacy.block_phone_number()
    if not _require_admin():
        return redirect(url_for("login.login"))

    phone = (request.form.get("phone") or "").strip()
    q = (request.form.get("q") or "").strip()
    status = (request.form.get("status") or "all").strip().lower()
    network = _normalize_network(request.form.get("network"))
    page = (request.form.get("page") or "1").strip()
    start_date = (request.form.get("start_date") or "").strip()
    end_date = (request.form.get("end_date") or "").strip()
    reason = (request.form.get("reason") or "").strip()

    key = _normalize_phone(phone)
    if re.fullmatch(r"0\d{9}", key):
        now = datetime.utcnow()
        blocked_phone_numbers_col.update_one(
            {"normalized_phone": key},
            {
                "$set": {
                    "phone": phone,
                    "normalized_phone": key,
                    "reason": reason,
                    "is_active": True,
                    "updated_at": now,
                    "blocked_by": session.get("admin_id") or session.get("user_id"),
                },
                "$setOnInsert": {"created_at": now},
            },
            upsert=True,
        )

    return redirect(
        url_for(
            "admin_phone_numbers.phone_numbers_page",
            q=q,
            status=status,
            network=network,
            start_date=start_date,
            end_date=end_date,
            page=page,
        )
    )


@admin_phone_numbers_bp.route("/admin/phone-numbers/unblock", methods=["POST"])
def unblock_phone_number():
    if session.get("role") == "admin":
        return legacy.unblock_phone_number()
    if not _require_admin():
        return redirect(url_for("login.login"))

    phone = (request.form.get("phone") or "").strip()
    q = (request.form.get("q") or "").strip()
    status = (request.form.get("status") or "all").strip().lower()
    network = _normalize_network(request.form.get("network"))
    page = (request.form.get("page") or "1").strip()
    start_date = (request.form.get("start_date") or "").strip()
    end_date = (request.form.get("end_date") or "").strip()

    key = _normalize_phone(phone)
    if re.fullmatch(r"0\d{9}", key):
        blocked_phone_numbers_col.update_one(
            {"normalized_phone": key, "is_active": True},
            {
                "$set": {
                    "is_active": False,
                    "updated_at": datetime.utcnow(),
                    "unblocked_by": session.get("admin_id") or session.get("user_id"),
                }
            },
        )

    return redirect(
        url_for(
            "admin_phone_numbers.phone_numbers_page",
            q=q,
            status=status,
            network=network,
            start_date=start_date,
            end_date=end_date,
            page=page,
        )
    )
